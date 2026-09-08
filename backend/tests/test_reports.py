"""
Sprint 7 tests: notifications, alert generation, report assembly,
PDF/Excel export, authorization, and error handling.
"""
import io
from openpyxl import load_workbook


def get_auth_headers(client, email="s7_creator@test.com"):
    client.post(
        "/api/auth/register",
        json={"name": "Sprint7 Creator", "email": email, "password": "test1234"},
    )
    login = client.post("/api/auth/login", json={"email": email, "password": "test1234"})
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def make_content(client, headers, **overrides):
    payload = {
        "platform": "youtube", "content_type": "video", "title": "S7 Video",
        "publish_date": "2026-01-15T10:00:00", "reach": 1000, "impressions": 1500,
        "likes": 80, "comments": 10, "shares": 5, "saves": 5, "views": 1200,
    }
    payload.update(overrides)
    return client.post("/api/content/", json=payload, headers=headers)


def make_revenue(client, headers, **overrides):
    payload = {
        "source": "youtube", "amount": 500.0, "currency": "USD", "date": "2026-01-15",
        "revenue_type": "ad_revenue", "status": "received",
    }
    payload.update(overrides)
    return client.post("/api/revenue/", json=payload, headers=headers)


# ---------- Notifications: CRUD ----------

def test_create_notification(client):
    headers = get_auth_headers(client)
    response = client.post("/api/notifications/", json={
        "notification_type": "performance",
        "title": "Test alert",
        "message": "Something happened.",
    }, headers=headers)
    assert response.status_code == 201
    assert response.json()["is_read"] is False


def test_list_notifications_includes_unread_count(client):
    headers = get_auth_headers(client)
    client.post("/api/notifications/", json={
        "notification_type": "revenue", "title": "A", "message": "msg",
    }, headers=headers)
    client.post("/api/notifications/", json={
        "notification_type": "revenue", "title": "B", "message": "msg",
    }, headers=headers)

    response = client.get("/api/notifications/", headers=headers)
    data = response.json()
    assert data["total"] == 2
    assert data["unread_count"] == 2


def test_mark_notification_read(client):
    headers = get_auth_headers(client)
    created = client.post("/api/notifications/", json={
        "notification_type": "engagement", "title": "A", "message": "msg",
    }, headers=headers).json()

    response = client.patch(
        f"/api/notifications/{created['id']}/read", json={"is_read": True}, headers=headers
    )
    assert response.status_code == 200
    assert response.json()["is_read"] is True

    listing = client.get("/api/notifications/", headers=headers).json()
    assert listing["unread_count"] == 0


def test_filter_notifications_by_type(client):
    headers = get_auth_headers(client)
    client.post("/api/notifications/", json={
        "notification_type": "performance", "title": "Perf", "message": "m",
    }, headers=headers)
    client.post("/api/notifications/", json={
        "notification_type": "revenue", "title": "Rev", "message": "m",
    }, headers=headers)

    response = client.get("/api/notifications/?notification_type=revenue", headers=headers)
    data = response.json()
    assert data["total"] == 1
    assert data["items"][0]["notification_type"] == "revenue"


def test_delete_notification(client):
    headers = get_auth_headers(client)
    created = client.post("/api/notifications/", json={
        "notification_type": "performance", "title": "A", "message": "m",
    }, headers=headers).json()

    delete_resp = client.delete(f"/api/notifications/{created['id']}", headers=headers)
    assert delete_resp.status_code == 204

    listing = client.get("/api/notifications/", headers=headers).json()
    assert listing["total"] == 0


def test_notification_not_found_returns_404(client):
    headers = get_auth_headers(client)
    fake_id = "00000000-0000-0000-0000-000000000000"
    response = client.patch(f"/api/notifications/{fake_id}/read", json={"is_read": True}, headers=headers)
    assert response.status_code == 404


def test_creator_cannot_access_another_creators_notification(client):
    headers_a = get_auth_headers(client, "s7_a@test.com")
    headers_b = get_auth_headers(client, "s7_b@test.com")

    created = client.post("/api/notifications/", json={
        "notification_type": "performance", "title": "Secret", "message": "m",
    }, headers=headers_a).json()

    response = client.delete(f"/api/notifications/{created['id']}", headers=headers_b)
    assert response.status_code == 404


def test_notifications_require_auth(client):
    response = client.get("/api/notifications/")
    assert response.status_code == 401


# ---------- Alert generation (reuses existing analytics services) ----------

def test_generate_alerts_creates_revenue_alert_for_pending(client):
    headers = get_auth_headers(client)
    make_revenue(client, headers, amount=500.0, status="pending")

    response = client.post("/api/notifications/generate", headers=headers)
    assert response.status_code == 201
    types = [n["notification_type"] for n in response.json()]
    assert "revenue" in types


def test_generate_alerts_creates_engagement_alert_for_low_rate(client):
    headers = get_auth_headers(client)
    make_content(client, headers, reach=1000, likes=1, comments=0, shares=0, saves=0)

    response = client.post("/api/notifications/generate", headers=headers)
    types = [n["notification_type"] for n in response.json()]
    assert "engagement" in types


def test_generate_alerts_creates_performance_alert_when_content_exists(client):
    headers = get_auth_headers(client)
    make_content(client, headers)

    response = client.post("/api/notifications/generate", headers=headers)
    types = [n["notification_type"] for n in response.json()]
    assert "performance" in types


def test_generate_alerts_with_no_data_creates_nothing_and_does_not_error(client):
    headers = get_auth_headers(client)
    response = client.post("/api/notifications/generate", headers=headers)
    assert response.status_code == 201
    assert response.json() == []


def test_generate_alerts_reuses_existing_service_not_recomputed(client):
    """Verifies the alert threshold check agrees with content_service's
    own engagement rate calculation -- proving reuse, not an independent
    recalculation that could silently drift out of sync."""
    headers = get_auth_headers(client)
    make_content(client, headers, reach=1000, likes=1, comments=0, shares=0, saves=0)

    kpi = client.get("/api/content/analytics/summary", headers=headers).json()
    alerts = client.post("/api/notifications/generate", headers=headers).json()

    engagement_alert = next(a for a in alerts if a["notification_type"] == "engagement")
    assert str(kpi["avg_engagement_rate"]) in engagement_alert["message"]


# ---------- Reports: JSON ----------

def test_get_creator_report_structure(client):
    headers = get_auth_headers(client)
    make_content(client, headers)
    make_revenue(client, headers)

    response = client.get("/api/reports/creator", headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert set(["creator", "content", "audience", "revenue", "growth", "platforms"]).issubset(data.keys())


def test_report_content_section_matches_content_analytics_endpoint(client):
    """Verifies the report's content KPI numbers are IDENTICAL to the
    existing Sprint 2 endpoint -- proving reuse, not reimplementation."""
    headers = get_auth_headers(client)
    make_content(client, headers, reach=1000, likes=80, comments=10, shares=5, saves=5)

    direct_kpi = client.get("/api/content/analytics/summary", headers=headers).json()
    report = client.get("/api/reports/creator", headers=headers).json()

    assert report["content"]["kpi_summary"] == direct_kpi


def test_report_revenue_section_matches_revenue_analytics_endpoint(client):
    headers = get_auth_headers(client)
    make_revenue(client, headers, amount=250.0)

    direct_kpi = client.get("/api/revenue/analytics/summary", headers=headers).json()
    report = client.get("/api/reports/creator", headers=headers).json()

    assert report["revenue"]["kpi_summary"] == direct_kpi


def test_report_with_no_data_returns_zeroed_sections_not_error(client):
    headers = get_auth_headers(client)
    response = client.get("/api/reports/creator", headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert data["content"]["kpi_summary"]["total_content"] == 0
    assert data["revenue"]["kpi_summary"]["total_revenue"] == 0.0


def test_reports_require_auth(client):
    response = client.get("/api/reports/creator")
    assert response.status_code == 401


def test_report_only_includes_requesting_creators_data(client):
    headers_a = get_auth_headers(client, "s7_report_a@test.com")
    headers_b = get_auth_headers(client, "s7_report_b@test.com")

    make_content(client, headers_a, title="Creator A Video")
    make_content(client, headers_b, title="Creator B Video")

    report_a = client.get("/api/reports/creator", headers=headers_a).json()
    titles_a = [c["title"] for c in report_a["content"]["top_performing"]]
    assert "Creator A Video" in titles_a
    assert "Creator B Video" not in titles_a


# ---------- Reports: PDF export ----------

def test_pdf_report_returns_valid_pdf_bytes(client):
    headers = get_auth_headers(client)
    make_content(client, headers)
    make_revenue(client, headers)

    response = client.get("/api/reports/creator/pdf", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content[:4] == b"%PDF"
    assert "attachment" in response.headers["content-disposition"]


def test_pdf_report_with_no_data_still_generates(client):
    headers = get_auth_headers(client)
    response = client.get("/api/reports/creator/pdf", headers=headers)
    assert response.status_code == 200
    assert response.content[:4] == b"%PDF"


def test_pdf_report_requires_auth(client):
    response = client.get("/api/reports/creator/pdf")
    assert response.status_code == 401


# ---------- Reports: Excel export ----------

def test_excel_report_returns_valid_workbook_with_expected_sheets(client):
    headers = get_auth_headers(client)
    make_content(client, headers)
    make_revenue(client, headers)

    response = client.get("/api/reports/creator/excel", headers=headers)
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(response.content))
    assert set(wb.sheetnames) == {
        "Summary", "Content Performance", "Audience Analytics",
        "Revenue", "Growth Trends", "Platform Comparison",
    }


def test_excel_report_values_match_backend_data_not_hardcoded(client):
    """Directly verifies a real number in the Excel file equals what
    was actually stored -- guards against hardcoded/placeholder values."""
    headers = get_auth_headers(client)
    make_revenue(client, headers, amount=333.33, date="2026-03-01")

    response = client.get("/api/reports/creator/excel", headers=headers)
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb["Revenue"]
    rows = list(ws.iter_rows(values_only=True))
    assert ("2026-03", 333.33) in rows


def test_excel_report_platform_values_are_plain_strings_not_enum_repr(client):
    """Regression test for a real bug: raw Platform enum objects were
    written directly to cells, producing 'Platform.youtube' instead of
    'youtube'."""
    headers = get_auth_headers(client)
    make_content(client, headers, platform="youtube")

    response = client.get("/api/reports/creator/excel", headers=headers)
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb["Content Performance"]
    rows = list(ws.iter_rows(values_only=True))
    platform_values = [row[1] for row in rows[1:]]
    assert "youtube" in platform_values
    assert not any(isinstance(v, str) and v.startswith("Platform.") for v in platform_values)


def test_excel_report_with_no_data_still_generates(client):
    headers = get_auth_headers(client)
    response = client.get("/api/reports/creator/excel", headers=headers)
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert "Summary" in wb.sheetnames


def test_excel_report_requires_auth(client):
    response = client.get("/api/reports/creator/excel")
    assert response.status_code == 401
