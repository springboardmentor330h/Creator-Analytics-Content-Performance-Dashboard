"""
Excel report generation. Same rule as pdf_report_service.py: this file
computes nothing, only lays out data already produced by
report_service.generate_creator_report() into sheets.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="7C5CFF", end_color="7C5CFF", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def generate_excel_report(report: dict) -> bytes:
    wb = Workbook()

    _build_summary_sheet(wb.active, report)
    _build_content_sheet(wb.create_sheet("Content Performance"), report)
    _build_audience_sheet(wb.create_sheet("Audience Analytics"), report)
    _build_revenue_sheet(wb.create_sheet("Revenue"), report)
    _build_growth_sheet(wb.create_sheet("Growth Trends"), report)
    _build_platform_sheet(wb.create_sheet("Platform Comparison"), report)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _enum_value(value) -> str:
    """
    Platform/ContentType/etc. are str-subclass Enums, so writing them
    directly to an openpyxl cell calls str() on them, which produces
    'Platform.youtube' instead of 'youtube' — Enum's __str__ is
    overridden even though the object IS a plain string underneath.
    Explicitly reading .value (or falling back to the value itself for
    plain strings) avoids that everywhere this helper is used.
    """
    return value.value if hasattr(value, "value") else value


def _write_header_row(ws, row_idx, headers):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left")


def _autosize_columns(ws, col_count, width=22):
    for i in range(1, col_count + 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _build_summary_sheet(ws, report):
    ws.title = "Summary"
    content_kpi = report["content"]["kpi_summary"]
    audience_kpi = report["audience"]["kpi_summary"]
    revenue_kpi = report["revenue"]["kpi_summary"]

    ws["A1"] = "CreatorIQ Analytics Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Creator: {report['creator']['name']} ({report['creator']['email']})"
    ws["A3"] = f"Generated: {report['generated_at'].strftime('%Y-%m-%d %H:%M UTC')}"

    _write_header_row(ws, 5, ["Metric", "Value"])
    rows = [
        ("Total Content", content_kpi["total_content"]),
        ("Avg Engagement Rate (%)", content_kpi["avg_engagement_rate"]),
        ("Total Followers", audience_kpi["total_followers"]),
        ("Audience Growth Rate (%)", audience_kpi["total_growth_rate_percent"]),
        ("Total Revenue ($)", revenue_kpi["total_revenue"]),
        ("Pending Revenue ($)", revenue_kpi["pending_revenue"]),
        ("Total Sponsorship Value ($)", revenue_kpi["total_sponsorship_value"]),
        ("Active Sponsorships", revenue_kpi["active_sponsorships"]),
    ]
    for i, (label, value) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)

    _autosize_columns(ws, 2, width=30)


def _build_content_sheet(ws, report):
    _write_header_row(ws, 1, ["Title", "Platform", "Content Type", "Reach", "Engagement Rate (%)"])
    for i, c in enumerate(report["content"]["top_performing"], start=2):
        ws.cell(row=i, column=1, value=c["title"])
        ws.cell(row=i, column=2, value=_enum_value(c["platform"]))
        ws.cell(row=i, column=3, value=_enum_value(c["content_type"]))
        ws.cell(row=i, column=4, value=c["reach"])
        ws.cell(row=i, column=5, value=c["engagement_rate"])
    _autosize_columns(ws, 5)


def _build_audience_sheet(ws, report):
    ws["A1"] = "Age Breakdown"
    ws["A1"].font = Font(bold=True)
    _write_header_row(ws, 2, ["Age Group", "Percentage"])
    row = 3
    for a in report["audience"]["age_breakdown"]:
        ws.cell(row=row, column=1, value=a["label"])
        ws.cell(row=row, column=2, value=a["percentage"])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Geographic Breakdown").font = Font(bold=True)
    row += 1
    _write_header_row(ws, row, ["Country", "Percentage"])
    row += 1
    for g in report["audience"]["geographic_breakdown"]:
        ws.cell(row=row, column=1, value=g["country"])
        ws.cell(row=row, column=2, value=g["percentage"])
        row += 1

    _autosize_columns(ws, 2)


def _build_revenue_sheet(ws, report):
    _write_header_row(ws, 1, ["Month", "Total Revenue ($)"])
    for i, t in enumerate(report["revenue"]["monthly_trend"], start=2):
        ws.cell(row=i, column=1, value=t["month"])
        ws.cell(row=i, column=2, value=t["total"])
    _autosize_columns(ws, 2)


def _build_growth_sheet(ws, report):
    _write_header_row(ws, 1, ["Platform", "Growth Rate (%)"])
    for i, g in enumerate(report["growth"]["by_platform"], start=2):
        ws.cell(row=i, column=1, value=_enum_value(g["platform"]))
        ws.cell(row=i, column=2, value=g["growth_rate_percent"])
    _autosize_columns(ws, 2)


def _build_platform_sheet(ws, report):
    _write_header_row(ws, 1, ["Platform", "Followers", "Content Count", "Avg Engagement (%)", "Growth (%)", "Data Source"])
    for i, s in enumerate(report["platforms"]["snapshots"], start=2):
        ws.cell(row=i, column=1, value=_enum_value(s["platform"]))
        ws.cell(row=i, column=2, value=s["followers"])
        ws.cell(row=i, column=3, value=s["total_content"])
        ws.cell(row=i, column=4, value=s["avg_engagement_rate"])
        ws.cell(row=i, column=5, value=s["growth_rate_percent"])
        ws.cell(row=i, column=6, value="Simulated" if s["is_mock_data"] else "Live")
    _autosize_columns(ws, 6)
