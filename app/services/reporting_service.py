from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy.orm import Session

from app.models.content import Content
from app.models.revenue import Revenue
from app.models.sponsorship import Sponsorship
from app.services.analytics_service import (
    kpi_summary,
    growth_report,
)
from app.services.audience_service import report as audience_report


def platform_comparison(
    db: Session,
    creator_id: UUID,
):
    rows = (
        db.query(Content)
        .filter(Content.creator_id == creator_id)
        .all()
    )

    result = {}

    for row in rows:
        platform = row.platform

        if platform not in result:
            result[platform] = {
                "content_count": 0,
                "views": 0,
                "likes": 0,
                "comments": 0,
                "shares": 0,
                "saves": 0,
                "reach": 0,
                "impressions": 0,
            }

        result[platform]["content_count"] += 1
        result[platform]["views"] += row.views or 0
        result[platform]["likes"] += row.likes or 0
        result[platform]["comments"] += row.comments or 0
        result[platform]["shares"] += row.shares or 0
        result[platform]["saves"] += row.saves or 0
        result[platform]["reach"] += row.reach or 0
        result[platform]["impressions"] += row.impressions or 0

    return result


def build_report(
    db: Session,
    creator_id: UUID,
):
    revenue_rows = (
        db.query(Revenue)
        .filter(Revenue.creator_uuid == creator_id)
        .all()
    )

    sponsorship_rows = (
        db.query(Sponsorship)
        .filter(Sponsorship.creator_uuid == creator_id)
        .all()
    )

    total_revenue = sum(
        row.amount for row in revenue_rows
    )

    revenue_by_source = {}

    for row in revenue_rows:
        revenue_by_source[row.source] = (
            revenue_by_source.get(row.source, 0)
            + row.amount
        )

    total_contract_value = sum(
        row.contract_value
        for row in sponsorship_rows
    )

    active_sponsorships = sum(
        1
        for row in sponsorship_rows
        if row.status == "active"
    )

    return {
        "creator_id": str(creator_id),

        "content": kpi_summary(
            db,
            creator_id=creator_id,
        ),

        "audience": audience_report(
            db,
            creator_id=creator_id,
        ),

        "growth": growth_report(
            db,
            days=30,
            creator_id=creator_id,
        ),

        "revenue": {
            "total_revenue": total_revenue,
            "revenue_by_source": revenue_by_source,
            "record_count": len(revenue_rows),
        },

        "sponsorships": {
            "total_contract_value": total_contract_value,
            "active_count": active_sponsorships,
            "total_count": len(sponsorship_rows),
        },

        "platform_comparison": platform_comparison(
            db,
            creator_id=creator_id,
        ),
    }


def pdf_report(data):
    buffer = BytesIO()

    pdf = canvas.Canvas(
        buffer,
        pagesize=A4,
    )

    width, height = A4
    y = height - 50

    pdf.setFont(
        "Helvetica-Bold",
        16,
    )

    pdf.drawString(
        50,
        y,
        "CreatorIQ Analytics Report",
    )

    y -= 30

    pdf.setFont(
        "Helvetica",
        9,
    )

    def write_value(text):
        nonlocal y

        if y < 50:
            pdf.showPage()
            y = height - 50
            pdf.setFont(
                "Helvetica",
                9,
            )

        text = str(text)

        if len(text) > 110:
            text = text[:107] + "..."

        pdf.drawString(
            50,
            y,
            text,
        )

        y -= 14

    def walk(obj, prefix=""):
        if isinstance(obj, dict):

            for key, value in obj.items():

                if isinstance(
                    value,
                    (dict, list),
                ):
                    write_value(
                        f"{prefix}{key}:"
                    )

                    walk(
                        value,
                        prefix + "  ",
                    )

                else:
                    write_value(
                        f"{prefix}{key}: {value}"
                    )

        elif isinstance(obj, list):

            for item in obj:

                if isinstance(
                    item,
                    dict,
                ):
                    walk(
                        item,
                        prefix,
                    )

                else:
                    write_value(
                        f"{prefix}{item}"
                    )

        else:
            write_value(
                f"{prefix}{obj}"
            )

    walk(data)

    pdf.save()

    buffer.seek(0)

    return buffer


def excel_report(data):
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"

    row = 1

    summary.cell(
        row,
        1,
        "CreatorIQ Analytics Report",
    )

    row += 2

    def write_object(
        sheet,
        obj,
        prefix="",
    ):
        nonlocal row

        if isinstance(obj, dict):

            for key, value in obj.items():

                if isinstance(
                    value,
                    (dict, list),
                ):

                    sheet.cell(
                        row,
                        1,
                        f"{prefix}{key}",
                    )

                    row += 1

                    write_object(
                        sheet,
                        value,
                        prefix + f"{key}.",
                    )

                else:

                    sheet.cell(
                        row,
                        1,
                        f"{prefix}{key}",
                    )

                    sheet.cell(
                        row,
                        2,
                        value,
                    )

                    row += 1

        elif isinstance(obj, list):

            for item in obj:

                if isinstance(
                    item,
                    dict,
                ):

                    write_object(
                        sheet,
                        item,
                        prefix,
                    )

                else:

                    sheet.cell(
                        row,
                        1,
                        prefix.rstrip("."),
                    )

                    sheet.cell(
                        row,
                        2,
                        str(item),
                    )

                    row += 1

        else:

            sheet.cell(
                row,
                1,
                prefix.rstrip("."),
            )

            sheet.cell(
                row,
                2,
                obj,
            )

            row += 1

    write_object(
        summary,
        data,
    )

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer