import io

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.services.analytics_service import get_dashboard_summary, get_platform_performance
from app.services.audience_service import get_audience_demographics
from app.services.revenue_service import get_revenue_summary


def build_creator_report_data(db: Session, creator_id: int):
    """Pulls together everything a combined report needs, reusing existing services."""
    return {
        "summary": get_dashboard_summary(db, creator_id),
"platform_performance": get_platform_performance(db, creator_id),
        "audience": get_audience_demographics(db, creator_id),
        "revenue": get_revenue_summary(db, creator_id),
    }


def generate_pdf_report(db: Session, creator_id: int) -> io.BytesIO:
    data = build_creator_report_data(db, creator_id)
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"CreatorIQ Performance Report — Creator #{creator_id}", styles["Title"]),
        Spacer(1, 16),
        Paragraph("Overview", styles["Heading2"]),
    ]

    summary = data["summary"]
    overview_rows = [["Metric", "Value"]] + [
        [key.replace("_", " ").title(), str(value)] for key, value in summary.items()
    ]
    elements.append(_styled_table(overview_rows))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Platform Performance", styles["Heading2"]))
    platform_rows = [["Platform", "Views", "Likes", "Avg. Engagement %"]]
    for p in data["platform_performance"]:
        platform_rows.append([
            p["platform"], str(p["total_views"]), str(p["total_likes"]),
            str(p["average_engagement_rate"]),
        ])
    elements.append(_styled_table(platform_rows))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Revenue Summary", styles["Heading2"]))
    revenue = data["revenue"]
    elements.append(Paragraph(f"Total revenue: {revenue['total_revenue']}", styles["Normal"]))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def _styled_table(rows) -> Table:
    table = Table(rows)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
    ]))
    return table


def generate_excel_report(db: Session, creator_id: int) -> io.BytesIO:
    data = build_creator_report_data(db, creator_id)
    wb = Workbook()

    summary_sheet = wb.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Metric", "Value"])
    for key, value in data["summary"].items():
        summary_sheet.append([key.replace("_", " ").title(), value])

    platform_sheet = wb.create_sheet("Platform Performance")
    platform_sheet.append(["Platform", "Total Views", "Total Likes", "Avg Engagement %"])
    for p in data["platform_performance"]:
        platform_sheet.append([
            p["platform"], p["total_views"], p["total_likes"], p["average_engagement_rate"],
        ])

    revenue_sheet = wb.create_sheet("Revenue")
    revenue_sheet.append(["Source", "Total"])
    for source, amount in data["revenue"]["by_source"].items():
        revenue_sheet.append([source, amount])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
